"""Maakt halloween-checkin-lijst.xlsx: één rij per speler, voor de check-in op zondag 1 november."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

F = 'Arial'
wb = Workbook()
ws = wb.active
ws.title = 'Check-in'

lime = PatternFill('solid', fgColor='CDFD50')
dark = PatternFill('solid', fgColor='0D181B')
lars = PatternFill('solid', fgColor='FFF2CC')    # Lars vult in
host = PatternFill('solid', fgColor='DDEBF7')    # check-in vult in
ex = PatternFill('solid', fgColor='EFEFEF')
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)

ws['A1'] = 'Halloween Padel Special · check-in · zondag 1 november 2026, 16:00 tot 23:00'
ws['A1'].font = Font(name=F, bold=True, size=14)
ws['A2'] = 'Geel = vult Lars in na de boeking (baannummer). Blauw = vult de check-in in op de dag. Wit = komt uit de betaling. Grijs = voorbeeld, mag weg.'
ws['A2'].font = Font(name=F, size=10, italic=True)
ws['M1'] = 'Totaal bij te betalen'
ws['M1'].font = Font(name=F, bold=True)
ws['N1'] = '=SUM(K5:K404)'
ws['N1'].number_format = '€ #,##0.00'
ws['N1'].font = Font(name=F, bold=True)
ws['M2'] = 'Spelers ingecheckt'
ws['M2'].font = Font(name=F, bold=True)
ws['N2'] = '=COUNTIF(I5:I404,"ja")'
ws['N2'].font = Font(name=F, bold=True)

head = ['Starttijd', 'Eindtijd', 'Baan', 'Boeker', 'Telefoon boeker', 'Speler', 'Duur (uur)',
        'Korting betaald (verkleed)', 'Aanwezig', 'Echt verkleed?', 'Bij te betalen (€)', 'Bijbetaald via', 'Opmerking', 'Mollie-betaling']
for c, h in enumerate(head, 1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.font = Font(name=F, bold=True, color='FFFFFF')
    cell.fill = dark
    cell.alignment = Alignment(wrap_text=True, vertical='center')
    cell.border = box
ws.row_dimensions[4].height = 32
ws['K4'].comment = Comment('Alleen als iemand korting heeft betaald maar niet echt verkleed is: de helft van zijn of haar deel. '
                           'Baanprijs €37,50 per uur (weekend, dubbelbaan, Playtomic), gedeeld door 4, daarvan de helft = €37,50 × duur / 8.', 'Lars')

example = [
    ('19:00', '20:30', 5, 'Lars de Bruin (voorbeeld)', '06 12345678', 'Lars de Bruin', 1.5, 'ja', 'ja', 'ja', None, None, None, 'tr_voorbeeld'),
    ('19:00', '20:30', 5, 'Lars de Bruin (voorbeeld)', '06 12345678', 'Anna', 1.5, 'ja', 'ja', 'ja', None, None, None, 'tr_voorbeeld'),
    ('19:00', '20:30', 5, 'Lars de Bruin (voorbeeld)', '06 12345678', 'Bram', 1.5, 'ja', 'ja', 'nee', None, 'Tikkie', 'alleen een hoedje', 'tr_voorbeeld'),
    ('19:00', '20:30', 5, 'Lars de Bruin (voorbeeld)', '06 12345678', 'Chris', 1.5, 'nee', 'ja', 'nee', None, None, 'heeft gewoon tarief betaald', 'tr_voorbeeld'),
]
FIRST, LAST = 5, 404
for r in range(FIRST, LAST + 1):
    vals = example[r - FIRST] if r - FIRST < len(example) else [None] * 14
    for c in range(1, 15):
        cell = ws.cell(row=r, column=c)
        if c != 11 and vals[c - 1] is not None:
            cell.value = vals[c - 1]
        cell.font = Font(name=F, size=10)
        cell.border = box
        if c == 3:
            cell.fill = lars
        elif c in (9, 10, 12, 13):
            cell.fill = host
        if r - FIRST < len(example) and c not in (3, 9, 10, 12, 13):
            cell.fill = ex
    ws.cell(row=r, column=11, value=f'=IF(AND(H{r}="ja",J{r}="nee"),ROUND(37.5*G{r}/8,2),0)').number_format = '€ #,##0.00;-€ #,##0.00;-'
    ws.cell(row=r, column=7).number_format = '0.0'

for col, opts in (('H', '"ja,nee"'), ('I', '"ja,nee"'), ('J', '"ja,nee"'), ('L', '"Tikkie,Pin,Contant"')):
    dv = DataValidation(type='list', formula1=opts, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'{col}{FIRST}:{col}{LAST}')

orange = PatternFill('solid', fgColor='F8CBAD')
ws.conditional_formatting.add(f'J{FIRST}:K{LAST}', FormulaRule(formula=[f'AND($H{FIRST}="ja",$J{FIRST}="nee")'], fill=orange))
greenf = PatternFill('solid', fgColor='C6EFCE')
ws.conditional_formatting.add(f'I{FIRST}:I{LAST}', FormulaRule(formula=[f'$I{FIRST}="ja"'], fill=greenf))

widths = [9, 9, 7, 24, 15, 20, 8, 12, 10, 11, 12, 13, 26, 16]
for i, w in enumerate(widths):
    ws.column_dimensions[chr(65 + i)].width = w
ws.freeze_panes = 'A5'
ws.auto_filter.ref = f'A4:N{LAST}'

u = wb.create_sheet('Uitleg check-in')
lines = [
    ('Halloween check-in: zo werkt het', True),
    ('Zondag 1 november 2026, 16:00 tot 23:00, bij Poort Padel. Spelers komen 10 minuten voor hun begintijd.', False),
    ('', False),
    ('1. Sorteer of filter op Starttijd. Elke boeking heeft 4 regels, één per speler.', False),
    ('2. Vraag de naam van de boeker en zoek de regels op.', False),
    ('3. Zet per speler Aanwezig op "ja".', False),
    ('4. Kijk per speler: echt verkleed of geschminkt? Zet Echt verkleed? op "ja" of "nee".', False),
    ('   Telt WEL: volledig kostuum, schmink of facepaint over het hele gezicht, masker met kostuum.', False),
    ('   Telt NIET: alleen een hoedje of oortjes, een gewoon shirt met print, twee streepjes schmink.', False),
    ('5. Wordt de cel oranje? Dan heeft die speler korting betaald maar is hij of zij niet echt verkleed.', False),
    ('   Bij te betalen staat in kolom K (1 uur: €4,69 · 1,5 uur: €7,03 · 2 uur: €9,38). Laat dat ter plekke bijbetalen en vul Bijbetaald via in.', False),
    ('6. Noem het baannummer (kolom Baan) en wens ze een griezelige wedstrijd.', False),
    ('', False),
    ('Twijfel over een kostuum? Kies voor de speler, het is een feestje. Mooie outfits: maak een foto voor Instagram (vraag eerst even of het mag).', False),
    ('Staat iemand niet op de lijst? Kijk in Mollie (Betalingen, zoek op naam) of in Netlify Forms (halloween-boeking). Niet betaald = gewoon tarief via Poort Padel.', False),
]
for i, (t, b) in enumerate(lines, 1):
    c = u.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=b, size=14 if b else 11)
u.column_dimensions['A'].width = 130

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
out = 'events/2026-halloween-padel/checkin/halloween-checkin-lijst.xlsx'
wb.save(out)
print(out)
