"""Maakt halloween-boekingen.xlsx: de sheet waar Make elke betaalde boeking in zet (tab Boekingen)
en een gesorteerde printlijst voor de bar van Poort Padel (tab Printlijst). Eén regel per boeking."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.properties import PageSetupProperties

F = 'Arial'
dark = PatternFill('solid', fgColor='0D181B')
lime = PatternFill('solid', fgColor='CDFD50')
lars = PatternFill('solid', fgColor='FFF2CC')
thin = Side(style='thin', color='999999')
box = Border(left=thin, right=thin, top=thin, bottom=thin)

COLS = ['Starttijd', 'Eindtijd', 'Baan', 'Boeker', 'Telefoon', 'Speler 1', 'Speler 2', 'Speler 3', 'Speler 4',
        'Verkleed', 'Betaald', 'Duur (uur)', 'Compleet', 'Kostuums ok', 'Bijbetaald', 'Opmerking', 'Mollie-id', 'Betaald op']
WIDTH = [8, 8, 6, 20, 13, 17, 17, 17, 17, 9, 9, 7, 9, 10, 10, 18, 16, 15]

wb = Workbook()
b = wb.active
b.title = 'Boekingen'
for c, h in enumerate(COLS, 1):
    x = b.cell(row=1, column=c, value=h)
    x.font = Font(name=F, bold=True, color='FFFFFF'); x.fill = dark; x.border = box
    x.alignment = Alignment(vertical='center', wrap_text=True)
b['C1'].fill = PatternFill('solid', fgColor='B7950B')
example = ['19:00', '20:30', '', 'VOORBEELD, deze regel mag weg', '06 12345678', 'Lars 🎃', 'Anna 🎃', 'Bram 🎃', 'Chris',
           '3 van 4', 35.16, 1.5, '', '', '', '', 'tr_voorbeeld', '17-10-2026 10:00']
for c, v in enumerate(example, 1):
    x = b.cell(row=2, column=c, value=v); x.font = Font(name=F, size=10, color='808080')
for r in range(2, 501):
    b.cell(row=r, column=11).number_format = '€ #,##0.00'
    b.cell(row=r, column=3).fill = lars
for i, w in enumerate(WIDTH):
    b.column_dimensions[chr(65 + i)].width = w
b.freeze_panes = 'A2'

p = wb.create_sheet('Printlijst')
p['A1'] = 'Halloween Padel Special · zaterdag 31 oktober · 16:00 tot 23:00 · check-inlijst Poort Padel'
p['A1'].font = Font(name=F, bold=True, size=14)
p['A2'] = ('🎃 = heeft korting betaald, moet echt verkleed of geschminkt zijn. Niet echt verkleed? Verschil bijbetalen: '
           '1 uur €4,69 · 1,5 uur €7,03 · 2 uur €9,38 per speler. Pas als de groep compleet is: baannummer noemen.')
p['A2'].font = Font(name=F, size=10, italic=True)
p.merge_cells('A2:P2')
p['A2'].alignment = Alignment(wrap_text=True, vertical='top')
p.row_dimensions[2].height = 30
for c, h in enumerate(COLS[:16], 1):
    x = p.cell(row=3, column=c, value=h)
    x.font = Font(name=F, bold=True, color='0D181B'); x.fill = lime; x.border = box
    x.alignment = Alignment(vertical='center', wrap_text=True)
p['A4'] = '=_xlfn._xlws.SORT(_xlfn._xlws.FILTER(Boekingen!A2:P500,Boekingen!A2:A500<>""),1,TRUE)'
for r in range(4, 120):
    for c in range(1, 17):
        x = p.cell(row=r, column=c); x.border = box; x.font = Font(name=F, size=10)
        x.alignment = Alignment(vertical='center')
    p.cell(row=r, column=11).number_format = '€ #,##0.00'
    p.cell(row=r, column=1).number_format = 'hh:mm'
    p.cell(row=r, column=2).number_format = 'hh:mm'
    p.row_dimensions[r].height = 24
for i, w in enumerate(WIDTH[:16]):
    p.column_dimensions[chr(65 + i)].width = w
p.page_setup.orientation = 'landscape'
p.page_setup.paperSize = p.PAPERSIZE_A4
p.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
p.page_setup.fitToWidth = 1
p.page_setup.fitToHeight = 0
p.print_title_rows = '3:3'
p.freeze_panes = 'A4'

out = 'events/2026-halloween-padel/checkin/halloween-boekingen.xlsx'
wb.save(out); print(out)
