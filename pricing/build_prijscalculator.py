"""Genereert pricing/prijscalculator.xlsx voor All Court Academy.

Prijsmodel: kostprijs (inkoop bij Poort Padel + coach-inzet) x (1 + winstmarge).
Alle formules zitten in het werkboek zelf, zodat Lars de inputs kan aanpassen.

Gebruik:  python3 pricing/build_prijscalculator.py
Daarna:   python3 <xlsx-skill>/scripts/recalc.py pricing/prijscalculator.xlsx
"""
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = "pricing/prijscalculator.xlsx"

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")          # invoer
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
GREEN = Font(name=FONT, color="008000")         # link naar ander blad
H1 = Font(name=FONT, bold=True, size=14)
YELLOW = PatternFill("solid", fgColor="FFFF00")  # aanname, nog te verifieren
GREY = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
EUR = '€#,##0.00;(€#,##0.00);"-"'
PCT = "0.0%"

wb = Workbook()

# ---------------------------------------------------------------- Inputs
ws = wb.active
ws.title = "Inputs"
ws["A1"] = "All Court Academy – prijscalculator events (inkoop Poort Padel + marge)"
ws["A1"].font = H1
ws["A2"] = "Blauw = invoer die je mag wijzigen. Geel = aanname, nog te verifiëren bij Poort Padel. Zwart = formule."
ws["A2"].font = Font(name=FONT, italic=True)

rows = [
    # key, label, value, unit, status, bron
    ("marge", "Winstmarge All Court Academy bovenop kostprijs", 0.20, "%", "Keuze Lars", "Opdracht: +20%"),
    ("btw", "BTW-tarief (sport/clinics 9%, horeca 9%, zaalhuur 21%) – info", 0.21, "%", "Info", "Belastingdienst; prijzen in dit bestand zijn excl. btw"),
    (None, "— Baanhuur dubbelbaan (Playtomic, geverifieerd 9 sep 2026) —", None, None, None, None),
    ("baan_dal", "Dubbelbaan daluren ma–vr 07:00–17:00, per uur", 30.00, "€/uur", "Geverifieerd", "playtomic.com/nl/clubs/poort-padel"),
    ("baan_piek", "Dubbelbaan piek ma–vr vanaf 17:00, per uur", 44.00, "€/uur", "Geverifieerd", "playtomic.com/nl/clubs/poort-padel"),
    ("baan_wknd", "Dubbelbaan weekend za/zo, per uur", 37.50, "€/uur", "Geverifieerd", "playtomic.com/nl/clubs/poort-padel"),
    ("spelers_per_baan", "Spelers per baan bij events", 4, "pers.", "Keuze Lars", "Standaard dubbelspel"),
    (None, "— Coaching (intern All Court Academy) —", None, None, None, None),
    ("coach_uur", "Kostprijs coach per uur (incl. reistijd/voorbereiding)", 45.00, "€/uur", "AANNAME", "Vul eigen coachtarief in"),
    ("spelers_per_coach", "Spelers per coach bij een clinic", 8, "pers.", "Keuze Lars", "2 banen per coach"),
    ("spelers_per_leider", "Spelers per toernooileider", 16, "pers.", "Keuze Lars", "4 banen per begeleider"),
    ("ballen_per_baan", "Ballen per baan per event", 2.00, "€/baan", "AANNAME", "Eigen inkoop ACA"),
    ("prijzen_event", "Prijsjes/medailles per toernooi", 25.00, "€/event", "AANNAME", "Eigen inkoop ACA"),
    ("racket_huur", "Rackethuur per persoon (0 = leenrackets ACA)", 0.00, "€/pers.", "AANNAME", "Poort Padel balieprijs niet online"),
    (None, "— Horeca Poort Padel (menukaart juni 2026) —", None, None, None, None),
    ("drankje", "Gemiddelde prijs per drankje (bier/fris/wijn)", 3.75, "€", "AANNAME", "Wijn €5–6,50 op de kaart; bier/fris niet online"),
    ("hapje", "Prijs per borrelhapje (teamschotel 25 st €24,50)", 0.98, "€", "Geverifieerd", "Menukaart PDF"),
    ("koffie", "Koffie/thee per kop", 3.00, "€", "AANNAME", "Niet op online kaart"),
    ("lunch_item", "Lunchgerecht gemiddeld (sandwich €10,50–14,50)", 12.50, "€", "Geverifieerd", "Menukaart PDF"),
    ("diner_item", "Dinergerecht gemiddeld (burger €17,50 / saté €19,50)", 18.50, "€", "Geverifieerd", "Menukaart PDF"),
    (None, "— Zaalhuur Poort Padel (NIET online; schatting) —", None, None, None, None),
    ("zaal_kantoor_uur", "Vergaderzaal/kantoor (10–12 pers.) per uur", 40.00, "€/uur", "AANNAME", "Offerte opvragen via reserveren.poortpadel.nl"),
    ("zaal_skybox_uur", "Skybox Londen/Parijs per uur", 50.00, "€/uur", "AANNAME", "Offerte opvragen"),
    ("zaal_congres_dagdeel", "Congreszaal Amsterdam (max 110) per dagdeel 4 uur", 400.00, "€/dagdeel", "AANNAME", "Offerte opvragen"),
    ("zaal_event_dagdeel", "Evenementenruimte (20–150) per dagdeel 4 uur", 500.00, "€/dagdeel", "AANNAME", "Offerte opvragen"),
    (None, "— Scholen —", None, None, None, None),
    ("leerlingen_per_baan", "Leerlingen per baan bij schoolclinic (roterend)", 6, "pers.", "Keuze Lars", "Bij 30 leerlingen = 5 banen"),
    ("afronding", "Verkoopprijs afronden op (bv. 0,50)", 0.50, "€", "Keuze Lars", ""),
]

hdr = ["Sleutel", "Omschrijving", "Waarde", "Eenheid", "Status", "Bron / toelichting"]
for c, h in enumerate(hdr, 1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.font = BOLD
    cell.fill = GREY
    cell.border = BOX

REF = {}  # key -> absolute reference "Inputs!$C$n"
r = 5
for key, label, value, unit, status, bron in rows:
    if key is None:
        ws.cell(row=r, column=2, value=label).font = BOLD
        r += 1
        continue
    ws.cell(row=r, column=1, value=key).font = Font(name=FONT, color="808080", size=9)
    ws.cell(row=r, column=2, value=label).font = BLACK
    v = ws.cell(row=r, column=3, value=value)
    v.font = BLUE
    v.number_format = PCT if unit == "%" else (EUR if unit and unit.startswith("€") else "0")
    if status == "AANNAME":
        v.fill = YELLOW
    ws.cell(row=r, column=4, value=unit).font = BLACK
    ws.cell(row=r, column=5, value=status).font = BLACK
    ws.cell(row=r, column=6, value=bron).font = Font(name=FONT, size=9)
    for c in range(1, 7):
        ws.cell(row=r, column=c).border = BOX
    REF[key] = f"Inputs!$C${r}"
    r += 1

ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 58
ws.column_dimensions["C"].width = 12
ws.column_dimensions["D"].width = 11
ws.column_dimensions["E"].width = 13
ws.column_dimensions["F"].width = 48
ws.freeze_panes = "A5"


def sell(cost_expr):
    """Verkoopprijs = kostprijs x (1+marge), afgerond naar boven op 'afronding'."""
    return f"=CEILING(({cost_expr})*(1+{REF['marge']}),{REF['afronding']})"


# ------------------------------------------------------- Arrangementen
wa = wb.create_sheet("Arrangementen")
wa["A1"] = "Zakelijke padelarrangementen – prijs per persoon (excl. btw, excl. horeca)"
wa["A1"].font = H1
wa["A2"] = "Minimum 16 personen (eis Poort Padel). Kostprijs = baanhuur/spelers + coach + toernooileider + prijzen + ballen. Verkoop = kostprijs x (1+marge)."
wa["A2"].font = Font(name=FONT, italic=True)

cols = [
    "Pakket", "Duur baan (uur)", "Clinic-uren coach", "Toernooi-uren leider", "Prijzen?",
    "Kostprijs daluren", "Kostprijs piek", "Kostprijs weekend",
    "Verkoop daluren", "Verkoop piek", "Verkoop weekend",
]
for c, h in enumerate(cols, 1):
    cell = wa.cell(row=4, column=c, value=h)
    cell.font = BOLD
    cell.fill = GREY
    cell.border = BOX
    cell.alignment = Alignment(wrap_text=True, vertical="center")
wa.row_dimensions[4].height = 32

pakketten = [
    ("A. Vrij spel (90 min)", 1.5, 0, 0, 0),
    ("B. Clinic + vrij spel (2 uur)", 2, 1, 0, 0),
    ("C. Toernooi (2 uur, Mexicano / King of the Court)", 2, 0, 2, 1),
    ("D. Clinic + toernooi (3 uur)", 3, 1, 2, 1),
    ("E. Padel Experience XL (4 uur: clinic, toernooi, vrij spel)", 4, 1, 2, 1),
]

r = 5
for naam, uren, clinic_u, toernooi_u, prijzen in pakketten:
    wa.cell(row=r, column=1, value=naam).font = BLACK
    for c, v in zip((2, 3, 4, 5), (uren, clinic_u, toernooi_u, prijzen)):
        cell = wa.cell(row=r, column=c, value=v)
        cell.font = BLUE
    # gedeelde kosten per persoon (coach, leider, prijzen, ballen, racket)
    shared = (
        f"C{r}*{REF['coach_uur']}/{REF['spelers_per_coach']}"
        f"+D{r}*{REF['coach_uur']}/{REF['spelers_per_leider']}"
        f"+E{r}*{REF['prijzen_event']}/{REF['spelers_per_leider']}"
        f"+{REF['ballen_per_baan']}/{REF['spelers_per_baan']}"
        f"+{REF['racket_huur']}"
    )
    for c, tarief in zip((6, 7, 8), ("baan_dal", "baan_piek", "baan_wknd")):
        wa.cell(row=r, column=c, value=f"=B{r}*{REF[tarief]}/{REF['spelers_per_baan']}+{shared}")
    for c, src in zip((9, 10, 11), ("F", "G", "H")):
        wa.cell(row=r, column=c, value=sell(f"{src}{r}"))
    for c in range(1, 12):
        cell = wa.cell(row=r, column=c)
        cell.border = BOX
        if c >= 6:
            cell.number_format = EUR
            cell.font = BLACK
    r += 1

wa.cell(row=r + 1, column=1, value="Toelichting").font = BOLD
notes = [
    "Daluren = ma–vr 07:00–17:00. Piek = ma–vr vanaf 17:00. Weekend = za/zo.",
    "Clinic-uren coach: uren dat een coach voor elke groep van 'spelers per coach' actief is.",
    "Toernooi-uren leider: uren toernooileiding per groep van 'spelers per toernooileider'.",
    "Prijzen? = 1 als er prijsjes/medailles worden uitgereikt.",
    "Horeca, zaalhuur en workshops komen hier bovenop (zie tabblad Add-ons).",
]
for i, n in enumerate(notes):
    wa.cell(row=r + 2 + i, column=1, value=n).font = Font(name=FONT, size=9)

wa.column_dimensions["A"].width = 52
for c in range(2, 12):
    wa.column_dimensions[get_column_letter(c)].width = 13
wa.freeze_panes = "B5"
ARR_FIRST_ROW = 5
ARR_LAST_ROW = 4 + len(pakketten)

# --------------------------------------------------------------- Add-ons
wd = wb.create_sheet("Add-ons")
wd["A1"] = "Add-ons per persoon of per zaal (excl. btw)"
wd["A1"].font = H1
hdr = ["Add-on", "Samenstelling", "Kostprijs", "Verkoopprijs", "Eenheid"]
for c, h in enumerate(hdr, 1):
    cell = wd.cell(row=3, column=c, value=h)
    cell.font = BOLD
    cell.fill = GREY
    cell.border = BOX

addons = [
    ("Koffie/thee bij ontvangst", "2 kopjes p.p.", f"=2*{REF['koffie']}", "p.p."),
    ("Borrelarrangement", "2 drankjes + 4 hapjes p.p.", f"=2*{REF['drankje']}+4*{REF['hapje']}", "p.p."),
    ("Borrelarrangement XL", "3 drankjes + 6 hapjes p.p.", f"=3*{REF['drankje']}+6*{REF['hapje']}", "p.p."),
    ("Lunch", "lunchgerecht + drankje + koffie", f"={REF['lunch_item']}+{REF['drankje']}+{REF['koffie']}", "p.p."),
    ("Diner", "dinergerecht + 2 drankjes", f"={REF['diner_item']}+2*{REF['drankje']}", "p.p."),
    ("Vergaderzaal/kantoor (10–12 pers.)", "dagdeel 4 uur", f"=4*{REF['zaal_kantoor_uur']}", "per zaal"),
    ("Skybox Londen/Parijs", "dagdeel 4 uur", f"=4*{REF['zaal_skybox_uur']}", "per zaal"),
    ("Congreszaal Amsterdam (max 110)", "dagdeel 4 uur", f"={REF['zaal_congres_dagdeel']}", "per zaal"),
    ("Evenementenruimte (20–150)", "dagdeel 4 uur", f"={REF['zaal_event_dagdeel']}", "per zaal"),
    ("Extra coach-uur (bv. clinic verlengen)", "1 uur coach", f"={REF['coach_uur']}", "per uur"),
    ("Extra baanuur daluren", "1 baan, 1 uur", f"={REF['baan_dal']}", "per baan"),
    ("Extra baanuur piek", "1 baan, 1 uur", f"={REF['baan_piek']}", "per baan"),
    ("Extra baanuur weekend", "1 baan, 1 uur", f"={REF['baan_wknd']}", "per baan"),
]
r = 4
for naam, samenstelling, kost, eenheid in addons:
    wd.cell(row=r, column=1, value=naam).font = BLACK
    wd.cell(row=r, column=2, value=samenstelling).font = BLACK
    k = wd.cell(row=r, column=3, value=kost)
    k.number_format = EUR
    v = wd.cell(row=r, column=4, value=sell(f"C{r}"))
    v.number_format = EUR
    wd.cell(row=r, column=5, value=eenheid).font = BLACK
    for c in range(1, 6):
        wd.cell(row=r, column=c).border = BOX
    r += 1
wd.cell(row=r + 1, column=1, value="Zaalhuur is een schatting tot Poort Padel de tarieven bevestigt (geel op tabblad Inputs).").font = Font(name=FONT, size=9)
wd.column_dimensions["A"].width = 40
wd.column_dimensions["B"].width = 32
wd.column_dimensions["C"].width = 13
wd.column_dimensions["D"].width = 14
wd.column_dimensions["E"].width = 10
ADDON_FIRST_ROW = 4
ADDON_LAST_ROW = 3 + len(addons)

# ---------------------------------------------------------------- Scholen
wsch = wb.create_sheet("Scholen")
wsch["A1"] = "Schoolarrangementen – daluren ma–vr onder schooltijd (excl. btw)"
wsch["A1"].font = H1
wsch["A2"] = "Kostprijs = banen x uren x daluren + coaches x uren x coachtarief + ballen. Verkoop = kostprijs x (1+marge)."
wsch["A2"].font = Font(name=FONT, italic=True)
hdr = ["Programma", "Leerlingen", "Banen", "Uren", "Coaches", "Kostprijs totaal", "Kostprijs p.l.", "Verkoop p.l.", "Verkoop totaal"]
for c, h in enumerate(hdr, 1):
    cell = wsch.cell(row=4, column=c, value=h)
    cell.font = BOLD
    cell.fill = GREY
    cell.border = BOX
    cell.alignment = Alignment(wrap_text=True, vertical="center")
wsch.row_dimensions[4].height = 30

scholen = [
    ("S1. Gymles op locatie – 1 klas, 60 min", 30, 1, 2),
    ("S2. Schoolclinic – 1 klas, 90 min", 30, 1.5, 2),
    ("S3. Sportdag / sportoriëntatie – 2 klassen, 3 uur roulerend", 60, 3, 3),
    ("S4. Naschools padelprogramma – 4 lessen x 60 min, 8 leerlingen (prijs per les)", 8, 1, 1),
    ("S5. Docenten-/teamuitje (zie Arrangementen B, daluren)", 16, 2, 2),
]
r = 5
for naam, ll, uren, coaches in scholen:
    wsch.cell(row=r, column=1, value=naam).font = BLACK
    wsch.cell(row=r, column=2, value=ll).font = BLUE
    if naam.startswith("S3"):
        # helft van de leerlingen tegelijk op de baan, andere helft doet nevenactiviteit
        wsch.cell(row=r, column=3, value=f"=CEILING(B{r}/2/{REF['leerlingen_per_baan']},1)")
    elif naam.startswith("S5"):
        wsch.cell(row=r, column=3, value=f"=CEILING(B{r}/{REF['spelers_per_baan']},1)")
    else:
        wsch.cell(row=r, column=3, value=f"=CEILING(B{r}/{REF['leerlingen_per_baan']},1)")
    wsch.cell(row=r, column=4, value=uren).font = BLUE
    wsch.cell(row=r, column=5, value=coaches).font = BLUE
    wsch.cell(row=r, column=6, value=f"=C{r}*D{r}*{REF['baan_dal']}+E{r}*D{r}*{REF['coach_uur']}+C{r}*{REF['ballen_per_baan']}")
    wsch.cell(row=r, column=7, value=f"=F{r}/B{r}")
    wsch.cell(row=r, column=8, value=sell(f"G{r}"))
    wsch.cell(row=r, column=9, value=f"=H{r}*B{r}")
    for c in range(1, 10):
        cell = wsch.cell(row=r, column=c)
        cell.border = BOX
        if c >= 6:
            cell.number_format = EUR
    r += 1
wsch.cell(row=r + 1, column=1, value="S4: prijs per les per leerling; een reeks van 4 lessen = 4x. Vervoer naar Poort Padel is voor de school (bushalte op 100 m).").font = Font(name=FONT, size=9)
wsch.column_dimensions["A"].width = 70
for c in range(2, 10):
    wsch.column_dimensions[get_column_letter(c)].width = 13
wsch.freeze_panes = "B5"

# ---------------------------------------------------------------- Offerte
wo = wb.create_sheet("Offerte")
wo["A1"] = "Snelle offerte – vul de blauwe cellen in"
wo["A1"].font = H1
labels = [
    ("Aantal deelnemers", 24, "pers."),
    ("Pakket (1–5: A, B, C, D, E)", 2, "nr"),
    ("Tijdvak (1 = daluren, 2 = piek, 3 = weekend)", 2, "nr"),
    ("Borrelarrangement erbij? (1 = ja)", 1, "ja/nee"),
    ("Lunch erbij? (1 = ja)", 0, "ja/nee"),
    ("Diner erbij? (1 = ja)", 0, "ja/nee"),
    ("Aantal vergaderzalen (dagdeel)", 0, "zalen"),
]
r = 3
for label, val, unit in labels:
    wo.cell(row=r, column=1, value=label).font = BLACK
    c = wo.cell(row=r, column=2, value=val)
    c.font = BLUE
    c.fill = YELLOW
    wo.cell(row=r, column=3, value=unit).font = BLACK
    r += 1
# r == 10 nu
wo.cell(row=11, column=1, value="Berekening").font = BOLD
calc = [
    ("Gekozen pakket", f"=INDEX(Arrangementen!$A${ARR_FIRST_ROW}:$A${ARR_LAST_ROW},B4)", None),
    ("Verkoopprijs pakket p.p.", f"=INDEX(Arrangementen!$I${ARR_FIRST_ROW}:$K${ARR_LAST_ROW},B4,B5)", EUR),
    ("Pakket totaal", "=B13*B3", EUR),
    ("Borrel p.p.", f"=B6*'Add-ons'!$D${ADDON_FIRST_ROW + 1}", EUR),
    ("Lunch p.p.", f"=B7*'Add-ons'!$D${ADDON_FIRST_ROW + 3}", EUR),
    ("Diner p.p.", f"=B8*'Add-ons'!$D${ADDON_FIRST_ROW + 4}", EUR),
    ("Horeca totaal", "=(B15+B16+B17)*B3", EUR),
    ("Zaalhuur totaal", f"=B9*'Add-ons'!$D${ADDON_FIRST_ROW + 5}", EUR),
    ("TOTAAL excl. btw", "=B14+B18+B19", EUR),
    ("Prijs per persoon excl. btw", "=IF(B3>0,B20/B3,0)", EUR),
    ("Waarvan marge All Court Academy (indicatie)", f"=B20-B20/(1+{REF['marge']})", EUR),
    ("Check: minimum 16 personen?", '=IF(B3>=16,"OK","LET OP: Poort Padel vraagt minimaal 16 personen")', None),
]
r = 12
for label, formula, fmt in calc:
    wo.cell(row=r, column=1, value=label).font = BOLD if label.startswith("TOTAAL") else BLACK
    c = wo.cell(row=r, column=2, value=formula)
    c.font = GREEN if "!" in formula else BLACK
    if fmt:
        c.number_format = fmt
    r += 1
wo.column_dimensions["A"].width = 48
wo.column_dimensions["B"].width = 52
wo.column_dimensions["C"].width = 10

# --------------------------------------------------------------- Legenda
wl = wb.create_sheet("Legenda")
legend = [
    ("Doel", "Prijslijst voor All Court Academy op basis van inkoop bij Poort Padel plus een winstmarge (standaard 20%)."),
    ("Blauw", "Invoercellen. Aanpassen mag."),
    ("Geel", "Aanname die nog geverifieerd moet worden bij Poort Padel (zaalhuur, horeca-arrangementen, rackethuur) of intern (coachtarief)."),
    ("Zwart", "Formules. Niet overschrijven."),
    ("Groen", "Verwijzing naar een ander tabblad."),
    ("Bron baanhuur", "Playtomic, club Poort Padel, tarieven bekeken op 6 en 9 september 2026."),
    ("Bron horeca", "Menukaart Poort Padel (PDF, juni 2026)."),
    ("Bron zalen", "reserveren.poortpadel.nl (AQQO) – capaciteiten; prijzen niet online."),
    ("Btw", "Alle bedragen excl. btw. Sportbeoefening/clinics vallen doorgaans onder 9%, zaalhuur onder 21%. Check met boekhouder."),
]
for i, (k, v) in enumerate(legend, 1):
    wl.cell(row=i, column=1, value=k).font = BOLD
    wl.cell(row=i, column=2, value=v).font = BLACK
wl.column_dimensions["A"].width = 16
wl.column_dimensions["B"].width = 110

# comments op aannames
ws[REF["coach_uur"].split("!")[1].replace("$", "")].comment = Comment(
    "Aanname. Vul het werkelijke uurtarief van je coaches in (loonkosten of zzp-tarief).", "Claude")
ws[REF["zaal_kantoor_uur"].split("!")[1].replace("$", "")].comment = Comment(
    "Schatting. Poort Padel publiceert geen zaalprijzen; vraag een offerte via reserveren.poortpadel.nl.", "Claude")

for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.font is None or cell.font.name != FONT:
                f = cell.font
                cell.font = Font(name=FONT, bold=f.bold, italic=f.italic, color=f.color, size=f.size or 11)

wb.save(OUT)
print("saved", OUT)
